"""Estimate whether text fits its box (no renderer needed).  Usage: python fitcheck.py deck.pptx
Measures with Arial (a bit wider than Calibri), so it is slightly pessimistic.
It cannot see overlaps between shapes or how a slide looks; open the deck to check that."""
import sys
from pptx import Presentation
from pptx.util import Emu
from PIL import ImageFont

REG = "/System/Library/Fonts/Supplemental/Arial.ttf"
BOLD = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
_cache = {}
def font(size, bold):
    key = (round(size * 4), bold)
    if key not in _cache:
        _cache[key] = ImageFont.truetype(BOLD if bold else REG, size=int(size * 4))   # 4x for precision
    return _cache[key]
def width_in(text, size, bold):
    return font(size, bold).getlength(text) / 4 / 72      # inches

def para_lines(words, size, bold, avail):
    lines, cur = 1, 0.0
    space = width_in(" ", size, bold)
    for w in words:
        ww = width_in(w, size, bold)
        if cur and cur + space + ww > avail:
            lines += 1; cur = ww
        else:
            cur = cur + (space if cur else 0) + ww
    return lines

def check(path):
    prs = Presentation(path)
    W, Hh = prs.slide_width, prs.slide_height
    problems = 0
    for n, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if sh.left is not None and (sh.left < -1 or sh.top < -1 or sh.left + sh.width > W + 9525 or sh.top + sh.height > Hh + 9525):
                print(f"slide {n}: '{sh.name}' extends past the slide edge"); problems += 1
            if not sh.has_text_frame or not sh.text_frame.text.strip() or sh.height == 0:
                continue   # (placeholders take size from the layout; not measured)
            tf = sh.text_frame
            bp = tf._txBody.bodyPr
            l = int(bp.get("lIns", 91440)); r = int(bp.get("rIns", 91440)); t = int(bp.get("tIns", 45720)); b = int(bp.get("bIns", 45720))
            avail_w = Emu(sh.width - l - r).inches
            avail_h = Emu(sh.height - t - b).inches
            need = 0.0
            for p in tf.paragraphs:
                runs = [x for x in p.runs if x.text]
                if not runs:
                    continue
                size = max((x.font.size.pt if x.font.size else (30 if sh.is_placeholder and n > 1 else 18)) for x in runs)
                bold = any(x.font.bold for x in runs)
                text = "".join(x.text for x in runs)
                has_bullet = p._p.pPr is not None and p._p.pPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar") is not None
                w = avail_w - (0.3 if has_bullet else 0)
                lines = para_lines(text.split(), size, bold, w)
                need += lines * size * 1.2 / 72
                spc = p._p.pPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}spcAft") if p._p.pPr is not None else None
                if spc is not None:
                    pts = spc.find("{http://schemas.openxmlformats.org/drawingml/2006/main}spcPts")
                    if pts is not None:
                        need += int(pts.get("val")) / 100 / 72
            if need > avail_h * 1.03:
                print(f"slide {n}: '{sh.name}' text needs {need:.2f} in but the box has {avail_h:.2f} in")
                problems += 1
    print("fit check:", "no problems found" if not problems else f"{problems} possible problem(s)")

check(sys.argv[1])
